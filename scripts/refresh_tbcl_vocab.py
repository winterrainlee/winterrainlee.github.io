from __future__ import annotations

import csv
import json
import re
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

SOURCE_URL = "https://coct.naer.edu.tw/file/files/14452%E8%A9%9E%E8%AA%9E%E8%A1%A8202504.xlsx"
SOURCE_VERSION = "TBCL 三等七級詞語表 2025-04"
OUT_DIR = Path("data/word-game")
XLSX_PATH = Path(".cache/tbcl-words-2025-04.xlsx")


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def normalize_header(value):
    return re.sub(r"\s+", "", clean(value))


def pick_col(headers, *needles):
    for idx, header in enumerate(headers):
        h = normalize_header(header)
        if any(n in h for n in needles):
            return idx
    return None


def parse_level(raw, fallback=""):
    text = clean(raw) or fallback
    text = text.replace("＊", "*")
    m = re.search(r"([1-7])\s*(\*)?", text)
    if not m:
        return ""
    return m.group(1) + ("*" if m.group(2) else "")


def detect_header(ws):
    for row_idx in range(1, min(ws.max_row, 30) + 1):
        vals = [clean(c.value) for c in ws[row_idx]]
        joined = "|".join(vals)
        if "詞語" in joined and "等級" in joined:
            return row_idx, vals
    return None, None


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    print(f"Downloading {SOURCE_URL}")
    urllib.request.urlretrieve(SOURCE_URL, XLSX_PATH)

    wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    rows = []

    for ws in wb.worksheets:
        header_row, headers = detect_header(ws)
        if not header_row:
            continue

        level_col = pick_col(headers, "等級")
        word_col = pick_col(headers, "詞語")
        category_col = pick_col(headers, "情境")
        zhuyin_col = pick_col(headers, "注音", "音讀")
        pinyin_col = pick_col(headers, "拼音")
        written_col = pick_col(headers, "書面")
        spoken_col = pick_col(headers, "口語")

        if level_col is None or word_col is None:
            continue

        last_level = ""
        for values in ws.iter_rows(min_row=header_row + 1, values_only=True):
            values = list(values)
            if level_col >= len(values) or word_col >= len(values):
                continue

            level = parse_level(values[level_col], last_level)
            if level:
                last_level = level
            if level not in {"4", "4*", "5"}:
                continue

            word = clean(values[word_col])
            if not word:
                continue

            def val(col):
                return clean(values[col]) if col is not None and col < len(values) else ""

            rows.append({
                "word": word,
                "tbcl_level": level,
                "cefr": "B1" if level.startswith("4") else "B2",
                "category": val(category_col),
                "zhuyin": val(zhuyin_col),
                "pinyin": val(pinyin_col),
                "written_freq_per_million": val(written_col),
                "spoken_freq_per_million": val(spoken_col),
            })

    if not rows:
        raise RuntimeError("No TBCL level 4/5 vocabulary rows were found. The source format may have changed.")

    # Preserve source order but remove accidental exact duplicates.
    seen = set()
    unique_rows = []
    for row in rows:
        key = (row["word"], row["tbcl_level"], row["zhuyin"], row["pinyin"])
        if key in seen:
            continue
        seen.add(key)
        unique_rows.append(row)
    rows = unique_rows

    counts = {
        "B1": sum(1 for r in rows if r["cefr"] == "B1"),
        "B2": sum(1 for r in rows if r["cefr"] == "B2"),
        "total": len(rows),
    }

    csv_path = OUT_DIR / "tbcl-b1-b2-master.csv"
    json_path = OUT_DIR / "tbcl-b1-b2-master.json"
    meta_path = OUT_DIR / "tbcl-b1-b2-source.json"

    fields = [
        "word", "tbcl_level", "cefr", "category", "zhuyin", "pinyin",
        "written_freq_per_million", "spoken_freq_per_million",
    ]
    with csv_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)

    with json_path.open("w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
        f.write("\n")

    metadata = {
        "source_name": SOURCE_VERSION,
        "source_url": SOURCE_URL,
        "source_authority": "國家教育研究院（NAER）華語文語料庫與能力基準整合應用系統",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "scope": "TBCL Level 4 / 4* / 5 (CEFR B1–B2 target pool)",
        "counts": counts,
        "notes": [
            "master 파일은 원자료 층이다. 한국어 뜻·게임 기믹·마을 배정 같은 편집 정보는 별도 overlay에 둔다.",
            "TBCL 4급은 B1, 5급은 B2 대응으로 사용한다.",
            "원자료의 공개 다운로드와 재배포 허용 범위는 별도 확인이 필요하다. 외부 공개 배포 전 라이선스를 재확인한다."
        ]
    }
    with meta_path.open("w", encoding="utf-8") as f:
        json.dump(metadata, f, ensure_ascii=False, indent=2)
        f.write("\n")

    print(json.dumps(counts, ensure_ascii=False))
    print(f"Wrote {csv_path}, {json_path}, {meta_path}")


if __name__ == "__main__":
    main()
